"""Relatório PDF dos valores salvos na tracking, sem macros nem instalação."""
from __future__ import annotations
import argparse
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
import os
import sys
from uuid import uuid4

PASTA = Path(__file__).resolve().parent
sys.path.insert(0, str(PASTA))
for wheel in sorted((PASTA/'bibliotecas').glob('*.whl')):
    sys.path.insert(0, str(wheel))
from openpyxl import load_workbook
from distribuir_cases import CABECALHOS, ErroBase, caminho_local, ler_identificador, ler_data
from painel_gestor import STATUS_VALIDOS


def ler_tracking(caminho):
    if not caminho.is_file() or caminho.suffix.lower() not in ('.xlsx','.xlsm'):
        raise ErroBase('Cole o caminho de uma tracking .xlsm ou .xlsx salva no computador.')
    livro = load_workbook(caminho, read_only=True, data_only=False, keep_links=False)
    try:
        if 'Tracking' not in livro.sheetnames:
            raise ErroBase('Não encontrei a aba Tracking. Use a planilha gerada pelo distribuidor.')
        ws = livro['Tracking']
        if [c.value for c in ws[1]][:7] != CABECALHOS:
            raise ErroBase('As sete colunas da aba Tracking foram alteradas. Confira o cabeçalho.')
        cases, vistos = [], set()
        for r, cells in enumerate(ws.iter_rows(min_row=2, max_col=7), 2):
            if all(c.value is None for c in cells):
                continue
            if any(c.data_type in ('f', 'e') for c in cells):
                raise ErroBase(f'Fórmula ou erro na linha {r} da Tracking. Preencha os campos com valores.')
            case = ler_identificador(cells[0], r)
            if case.casefold() in vistos:
                raise ErroBase(f'CASE duplicado na tracking, linha {r}. Corrija antes de gerar o PDF.')
            vistos.add(case.casefold())
            abertura = ler_data(cells[1].value, livro.epoch, 'Data de abertura', r)
            # Mesma semântica do Excel: status fora da lista/espacos extras são sinalizados.
            raw = str(cells[3].value or '')
            status = next((s for s in STATUS_VALIDOS if s.casefold() == raw.casefold()), 'A corrigir')
            cases.append({'case':case, 'abertura':abertura, 'status':status,
                          'responsavel':str(cells[6].value or ''), 'observacao':str(cells[5].value or '')})
        nomes = []
        origem = None
        recebidos = excluidos = None
        if 'Resumo' in livro.sheetnames:
            res = livro['Resumo']
            origem = str(res['F6'].value or '')
            recebidos, excluidos = res['F9'].value, res['F10'].value
            qtd = res['F12'].value
            if isinstance(qtd, int) and 0 <= qtd <= 1000:
                nomes = [str(res.cell(r,2).value or '') for r in range(19,19+qtd)]
        por_pessoa = defaultdict(Counter)
        status_total = Counter()
        cadastro = {n.casefold(): n for n in nomes}
        sem_cadastro = 0
        for c in cases:
            nome = cadastro.get(c['responsavel'].casefold())
            if nome is None:
                sem_cadastro += 1
                nome = c['responsavel'] or '(sem responsável)'
                if nome not in nomes:
                    nomes.append(nome)
            por_pessoa[nome][c['status']] += 1
            por_pessoa[nome]['Total'] += 1
            status_total[c['status']] += 1
        return {'cases':cases, 'nomes':nomes, 'por_pessoa':por_pessoa,
                'status':status_total, 'total':len(cases), 'sem_cadastro':sem_cadastro,
                'recebidos':recebidos, 'excluidos':excluidos, 'origem':origem}
    finally:
        livro.close()


def gerar_pdf(caminho, pasta=None):
    from fpdf import FPDF
    import fpdf.fpdf as engine
    engine.FPDF_CACHE_MODE = 1  # Nunca gravar cache na pasta de fontes do Windows.
    dados = ler_tracking(caminho)
    agora = datetime.now()
    pasta = pasta or caminho.parent/'Relatorios_PDF'
    pasta.mkdir(parents=True, exist_ok=True)
    arquivo = pasta/('Resumo_Gestor_'+agora.strftime('%d-%m_%H-%M-%S_')+uuid4().hex[:6]+'.pdf')
    fonts = Path(os.environ.get('WINDIR', 'C:/Windows'))/'Fonts'
    pdf = FPDF(orientation='L', unit='mm', format='A4')
    pdf.set_auto_page_break(False)
    pdf.set_margins(10,10,10)
    # Fonte Unicode local já presente no Windows; não exige rede nem instalação.
    if not (fonts/'arial.ttf').is_file() or not (fonts/'arialbd.ttf').is_file():
        raise ErroBase('As fontes Arial do Windows não foram encontradas. Exporte a aba Gestor pelo Excel (Ctrl+P).')
    pdf.add_font('Tracking','',str(fonts/'arial.ttf'),uni=True)
    pdf.add_font('Tracking','B',str(fonts/'arialbd.ttf'),uni=True)
    BG=(16,25,35); CARD=(28,41,56); WHITE=(242,246,250); MUTED=(173,190,208)
    TEAL=(115,222,201); BLUE=(138,175,255); AMBER=(245,203,131); PURPLE=(198,172,250)

    def text(x,y,s,size=10,color=WHITE,bold=False):
        pdf.set_font('Tracking','B' if bold else '',size)
        pdf.set_text_color(*color)
        pdf.set_xy(x,y)
        pdf.cell(0,5,str(s))

    def wrap(s,width,size=10):
        pdf.set_font('Tracking','',size)
        lines=[]
        for para in str(s).splitlines() or ['']:
            line=''
            for word in para.split():
                if line and pdf.get_string_width(line+' '+word)>width:
                    lines.append(line);line=''
                # Também quebra identificadores/observações sem espaços.
                while pdf.get_string_width(word)>width:
                    cut=1
                    while cut<len(word) and pdf.get_string_width(word[:cut+1])<=width:
                        cut+=1
                    if line: lines.append(line);line=''
                    lines.append(word[:cut]);word=word[cut:]
                line=(line+' '+word).strip()
            lines.append(line)
        return lines

    def pagina(titulo):
        pdf.add_page()
        pdf.set_fill_color(*BG);pdf.rect(0,0,297,210,'F')
        text(10,9,'TRACKING DE CASES',10,TEAL,True)
        text(10,19,titulo,20,WHITE,True)
        text(10,31,f'PDF gerado em {agora:%d/%m/%Y %H:%M:%S} | Arquivo salvo em {datetime.fromtimestamp(caminho.stat().st_mtime):%d/%m/%Y %H:%M:%S}',9,MUTED)
        text(10,198,'Retrato dos dados salvos. Escalados e revisão permanecem abertos.',9,MUTED)
        text(272,198,f'{pdf.page_no():02}',9,MUTED)
        return 44

    def table_header(y,labels,widths):
        pdf.set_fill_color(*CARD);pdf.rect(10,y,277,9,'F')
        x=12
        for label,w in zip(labels,widths):
            text(x,y+2,label,9,MUTED,True);x+=w
        return y+10

    y=pagina('Acompanhamento da operação')
    stats=dados['status'];total=dados['total'];closed=stats['Concluído']
    for x,label,val,col in [(10,'Distribuídos',total,BLUE),(104,'Concluídos',closed,TEAL),(198,'Ainda abertos',total-closed,AMBER)]:
        pdf.set_fill_color(*CARD);pdf.rect(x,y,89,27,'F')
        text(x+5,y+3,label,10,MUTED);text(x+5,y+13,val,23,col,True)
    y+=32
    for x,label,val,col in [(10,'Pendentes',stats['Pendente'],MUTED),(81,'Em andamento',stats['Em andamento'],BLUE),(152,'Revisar depois',stats['Revisar depois'],AMBER),(223,'Escalados',stats['Escalado'],PURPLE)]:
        text(x,y,label,10,MUTED);text(x,y+7,val,17,col,True)
    y+=23
    labels=['Responsável','Total','Concluídos','Abertos','Pendentes','Andamento','Revisão','Escalados']
    widths=[67,30,30,30,30,30,30,30]
    y=table_header(y,labels,widths)
    for nome in dados['nomes']:
        ct=dados['por_pessoa'][nome]
        linhas=wrap(nome,61)
        h=max(9,5*len(linhas)+3)
        if y+h>179:
            y=pagina('Acompanhamento por responsável');y=table_header(y,labels,widths)
        for k,line in enumerate(linhas):text(12,y+1+k*5,line)
        vals=[ct['Total'],ct['Concluído'],ct['Total']-ct['Concluído'],ct['Pendente'],ct['Em andamento'],ct['Revisar depois'],ct['Escalado']]
        x=79
        for val in vals:text(x,y+1,val,10,TEAL);x+=30
        y+=h
    if y+17>189:y=pagina('Conferência do acompanhamento')
    text(10,y+3,f"Status a corrigir: {stats['A corrigir']}  |  Sem responsável cadastrado: {dados['sem_cadastro']}  |  Conclusão: {closed/total:.0%}" if total else 'Nenhum case na tracking.',10,AMBER)
    if dados['recebidos'] is not None:
        text(10,y+10,f"Distribuição inicial: {dados['recebidos']} recebidos | {dados['excluidos']} excluídos em aguardo de chamado.",9,MUTED)

    escalados=[c for c in dados['cases'] if c['status']=='Escalado']
    if escalados:
        y=pagina('Cases escalados | destino e contexto')
        for c in escalados:
            lines=wrap(f"Case {c['case']} | Responsável: {c['responsavel'] or '(não informado)'}",267,11)
            lines+=wrap(c['observacao'] or 'Observação não preenchida: informe equipe/pessoa de destino e contexto.',267,10)
            for i,line in enumerate(lines):
                if y>185:y=pagina('Cases escalados | continuação')
                text(12,y,line,11 if i==0 else 10,TEAL if i==0 else WHITE,i==0);y+=6
            y+=5
    pdf.set_title('Tracking de cases - Relatório do gestor')
    pdf.set_author('Tracking operacional')
    pdf.output(str(arquivo),'F')
    return arquivo


def main():
    parser=argparse.ArgumentParser(description='Gera o PDF do acompanhamento salvo na tracking.')
    parser.add_argument('arquivo',nargs='?')
    parser.add_argument('--saida',type=Path)
    parser.add_argument('--sem-pausa',action='store_true')
    args=parser.parse_args();codigo=0
    try:
        print('\nRELATÓRIO DO GESTOR | PDF\nSalve a tracking no Excel antes de continuar. O PDF usa os dados salvos.\n')
        arquivo=caminho_local(args.arquivo or input('Cole o caminho completo da tracking .xlsm ou .xlsx: '))
        resultado=gerar_pdf(arquivo,args.saida)
        print(f'\nPDF gerado:\n{resultado.resolve()}')
    except (KeyboardInterrupt,EOFError):
        codigo=130;print('\nOperação cancelada.')
    except Exception as exc:
        codigo=1;print(f'\nNão foi possível gerar o PDF: {exc}')
    if not args.sem_pausa:
        try:input('\nPressione Enter para fechar...')
        except (KeyboardInterrupt,EOFError):pass
    return codigo


if __name__=='__main__':
    raise SystemExit(main())
