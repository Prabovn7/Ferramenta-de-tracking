"""SLA de quatro dias úteis, encerrado ao fim do quarto dia após a abertura."""
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.formatting.rule import FormulaRule
from painel_gestor import texto, fundo, BG, CARD, TEXT, MUTED, TEAL, AMBER, BLUE

RED = 'FF9098'


def acrescentar_sla(livro, gestor, inicio):
    # Preserva feriados ao atualizar uma planilha que já possui o SLA.
    feriados = []
    if 'SLA' in livro:
        antigo = livro['SLA']
        feriados = [antigo.cell(r, 11).value for r in range(10, 110)]
        del livro['SLA']
    sla = livro.create_sheet('SLA')
    sla.sheet_properties.tabColor = AMBER
    quantidade = livro['Tracking'].max_row - 1
    fim = 7 + quantidade
    fundo(sla, max(109, fim), 12)
    for col, width in {'A':17,'B':27,'C':27,'D':20,'E':26,'F':18,'G':16,'H':27,'I':3,'J':23,'K':23,'L':3}.items():
        sla.column_dimensions[col].width = width
    sla.freeze_panes = 'C8'
    sla.sheet_view.zoomScale = 80
    sla.print_area = f'A1:H{fim}'
    sla.print_title_rows = '7:7'
    texto(sla, 'A1:H2', 'SLA por case', 24, TEXT, True)
    texto(sla, 'A3:H3', 'Quatro dias úteis após a abertura. O dia de abertura não entra na contagem.', 11, MUTED)
    texto(sla, 'A4:H4', 'Prazo até 23:59:59. Sábados, domingos e feriados cadastrados não contam.', 11, MUTED)
    texto(sla, 'A5:H5', 'Atualize os atendimentos na Tracking. Filtre esta lista por responsável ou situação.', 11, TEAL)
    texto(sla, 'A6:H6', 'Concluídos saem do SLA em aberto; não há data de conclusão para avaliar o SLA histórico.', 10, MUTED)
    texto(sla, 'J2:K2', 'REGRA DO SLA', 12, TEAL, True)
    texto(sla, 'J3', 'Prazo em dias úteis', 11, MUTED)
    texto(sla, 'K3', 4, 14, TEXT, True)
    texto(sla, 'J5:K5', 'Data e hora do cálculo', 11, MUTED)
    texto(sla, 'J6:K6', "='Gestor'!$Z$2", 12, TEXT, True).number_format = 'dd/mm/yyyy hh:mm:ss'
    texto(sla, 'J8:K8', 'Feriados: informe datas na coluna K', 11, AMBER, True)
    texto(sla, 'J9:K9', '=IF(COUNT(K10:K109)=0,"Sem feriados cadastrados",COUNT(K10:K109)&" datas cadastradas")', 10, MUTED)
    for r in range(10, 110):
        sla.cell(r, 11).fill = PatternFill('solid', fgColor='324459')
        sla.cell(r, 11).number_format = 'dd/mm/yyyy'
        if r - 10 < len(feriados):
            sla.cell(r, 11, feriados[r - 10])
    validacao = DataValidation(type='date', operator='between', formula1='DATE(1900,1,1)', formula2='DATE(2100,12,31)', allow_blank=True)
    validacao.showErrorMessage = True
    validacao.errorTitle = 'Informe uma data válida'
    validacao.error = 'Use uma data do Excel, como 12/10/2026.'
    sla.add_data_validation(validacao)
    validacao.add('K10:K109')
    headers = ['CASE','DATA DE ABERTURA','RESPONSÁVEL','STATUS','PRAZO ATÉ','DIAS ÚTEIS','FAIXA SLA','SITUAÇÃO DO SLA']
    for c, h in enumerate(headers, 1):
        cell = sla.cell(7, c, h)
        cell.fill = PatternFill('solid', fgColor='263B4F')
        cell.font = Font(name='Arial', size=11, color=TEXT, bold=True)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    sla.row_dimensions[7].height = 32
    fer = '$K$10:$K$109'
    for r in range(8, fim + 1):
        source = r - 6
        for col, orig in [('A','A'),('B','B'),('C','G'),('D','D')]:
            sla[f'{col}{r}'] = f'=IF(\'Tracking\'!{orig}{source}="","",\'Tracking\'!{orig}{source})'
        sla[f'E{r}'] = f'=IF(A{r}="","",IF(ISNUMBER(B{r}),IF(B{r}>0,WORKDAY(INT(B{r}),$K$3,{fer})+1-1/86400,"Verificar data"),"Verificar data"))'
        sla[f'H{r}'] = f'=IF(A{r}="","",IF(D{r}="Concluído","Concluído",IF(ISNUMBER(E{r}),IF(B{r}>$J$6,"Data inválida",IF(INT($J$6)>INT(E{r}),"Vencido",IF(INT($J$6)=INT(E{r}),"Vence hoje","Dentro do prazo"))),"Data inválida")))'
        sla[f'F{r}'] = f'=IF(OR(A{r}="",D{r}="Concluído",H{r}="Data inválida"),"",MAX(0,NETWORKDAYS(INT(B{r})+1,INT($J$6),{fer})))'
        sla[f'G{r}'] = f'=IF(A{r}="","",IF(D{r}="Concluído","Concluído",IF(H{r}="Data inválida","Verificar data",IF(H{r}="Vencido","Vencido","D"&F{r}))))'
        for c in range(1, 9):
            sla.cell(r, c).font = Font(name='Arial', size=11, color=TEXT)
            sla.cell(r, c).fill = PatternFill('solid', fgColor=CARD if r % 2 == 0 else BG)
        sla[f'A{r}'].number_format = '00000000'
        sla[f'B{r}'].number_format = sla[f'E{r}'].number_format = 'dd/mm/yyyy hh:mm:ss'
        sla[f'F{r}'].number_format = '0'
    if quantidade:
        tab = Table(displayName='DetalheSLA', ref=f'A7:H{fim}')
        tab.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
        sla.add_table(tab)
        for status, color, fill in [('Vencido', RED, '472A35'),('Vence hoje', AMBER, '483C27'),('Dentro do prazo', TEAL, CARD)]:
            sla.conditional_formatting.add(f'G8:H{fim}', FormulaRule(formula=[f'$H8="{status}"'], font=Font(name='Arial', color=color, bold=True), fill=PatternFill('solid', fgColor=fill)))

    # Limites comuns calculados no Gestor: o botão existente recalcula esta aba
    # antes de exportar. As contagens não dependem de caches da aba de detalhes.
    fer = "'SLA'!$K$10:$K$109"
    helper = {
        2: ('Instante do cálculo','=NOW()'),
        3: ('Hoje é dia útil', f'=NETWORKDAYS(INT(Z2),INT(Z2),{fer})>0'),
        4: ('Último dia útil', f'=WORKDAY(INT(Z2)+1,-1,{fer})'),
        5: ('Abertura: limite exclusivo dos vencidos', f'=IF(Z3,WORKDAY(INT(Z2),-\'SLA\'!$K$3,{fer}),WORKDAY(Z4,1-\'SLA\'!$K$3,{fer}))'),
        6: ('Abertura: início dos prazos futuros', f'=WORKDAY(Z4,1-\'SLA\'!$K$3,{fer})'),
    }
    for r, (label, formula) in helper.items():
        gestor.cell(r, 25, label)
        gestor.cell(r, 26, formula)
    gestor.column_dimensions['Y'].hidden = True
    gestor.column_dimensions['Z'].hidden = True
    ids = 'CasesDistribuidos[CASE]'
    dates = 'CasesDistribuidos[DATA DE ABERTURA]'
    sts = 'CasesDistribuidos[STATUS]'
    owners = 'CasesDistribuidos[RESPONSÁVEL]'
    criterio = 'SUBSTITUTE(SUBSTITUTE(SUBSTITUTE($E$5,"~","~~"),"*","~*"),"?","~?")'
    def count(extra):
        filters = f'{ids},"<>",{sts},"<>Concluído",{dates},">0",{dates},"<="&$Z$2,{extra}'
        return f'IF($E$5="Toda a equipe",COUNTIFS({filters}),COUNTIFS({filters},{owners},{criterio}))'
    texto(gestor, f'B{inicio}:V{inicio}', '="SLA DOS CASES EM ABERTO  /  "&$E$5', 12, TEAL, True)
    texto(gestor, f'B{inicio+1}:V{inicio+1}', '4 dias úteis: até o fim do quarto dia após a abertura. Inclui escalados e cases para revisão.', 10, MUTED)
    specs = [
        ('B','G','PRAZO APÓS HOJE','='+count(f'{dates},">="&$Z$6'),TEAL),
        ('I','N','VENCEM HOJE','=IF($Z$3,'+count(f'{dates},">="&$Z$5,{dates},"<"&$Z$6')+',0)',AMBER),
        ('P','V','SLA VENCIDO','='+count(f'{dates},"<"&$Z$5'),RED),
    ]
    for left, right, label, formula, color in specs:
        texto(gestor, f'{left}{inicio+3}:{right}{inicio+3}', label, 11, MUTED, True, CARD)
        texto(gestor, f'{left}{inicio+4}:{right}{inicio+5}', formula, 26, color, True, CARD).number_format = '#,##0'
    texto(gestor, f'B{inicio+6}:K{inicio+6}', 'Datas inválidas ou futuras em abertos', 10, MUTED)
    texto(gestor, f'L{inicio+6}:M{inicio+6}', f'=P8-SUM(B{inicio+4},I{inicio+4},P{inicio+4})', 12, AMBER, True)
    texto(gestor, f'O{inicio+6}:V{inicio+6}', '=IF(COUNT(\'SLA\'!$K$10:$K$109)=0,"Sem feriados cadastrados","Feriados cadastrados: "&COUNT(\'SLA\'!$K$10:$K$109))', 10, MUTED)
    texto(gestor, f'B{inicio+7}:K{inicio+7}', 'Calculado em: abra, edite ou pressione F9', 10, MUTED)
    texto(gestor, f'L{inicio+7}:V{inicio+7}', '=$Z$2', 11, TEXT).number_format = 'dd/mm/yyyy hh:mm:ss'
    link = texto(gestor, f'B{inicio+9}:M{inicio+9}', 'VER PRAZO E FAIXA SLA DE CADA CASE', 11, TEAL, True, CARD)
    link.hyperlink = "#'SLA'!A1"
    link = texto(gestor, f'O{inicio+9}:V{inicio+9}', 'CADASTRAR FERIADOS', 11, AMBER, True, CARD)
    link.hyperlink = "#'SLA'!K10"
    gestor['L5'] = 'VER SLA DE 4 DIAS ÚTEIS'
    gestor['L5'].font = Font(name='Arial', size=11, color=TEAL, bold=True)
    gestor['L5'].hyperlink = f"#'Gestor'!B{inicio}"
