"""Atualiza o painel da versão 2 sem redistribuir nem zerar o atendimento."""
from pathlib import Path
from datetime import datetime
from uuid import uuid4
import argparse
import sys
sys.path.insert(0,str(Path(__file__).resolve().parent))
from distribuir_cases import load_workbook, ErroBase, caminho_local, CABECALHOS
from openpyxl.workbook.properties import CalcProperties
from dashboard_gestor import desenhar
from incluir_botao_pdf import incluir_botao


def atualizar(origem, pasta=None):
    if origem.suffix.lower()!='.xlsx' or not origem.is_file():
        raise ErroBase('Use o arquivo .xlsx da versão 2, com Tracking e Resumo.')
    w=load_workbook(origem,keep_links=False)
    tmp=None
    try:
        if not {'Tracking','Resumo'}.issubset(w.sheetnames):
            raise ErroBase('A planilha precisa ter as abas Tracking e Resumo da versão 2.')
        if [c.value for c in w['Tracking'][1]]!=CABECALHOS:
            raise ErroBase('O cabeçalho da Tracking foi alterado. Confira as sete colunas.')
        resumo=w['Resumo'];qtd=resumo['F12'].value
        if not isinstance(qtd,int) or not 1<=qtd<=1000:
            raise ErroBase('Não consegui ler a equipe na aba Resumo.')
        nomes=[resumo.cell(r,2).value for r in range(19,19+qtd)]
        if any(not isinstance(n,str) or not n for n in nomes):
            raise ErroBase('Há nomes vazios no Resumo. Confira a lista da equipe.')
        if 'Gestor' in w:del w['Gestor']
        g=w.create_sheet('Gestor',0);g.sheet_properties.tabColor='73DEC9'
        desenhar(g,resumo,nomes,w)
        w.active=0
        w.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True,forceFullCalc=True,calcMode='auto')
        pasta=pasta or origem.parent/('Painel_v3_'+datetime.now().strftime('%Y-%m-%d_%H-%M-%S_')+uuid4().hex[:6])
        pasta.mkdir(parents=True,exist_ok=True)
        destino=pasta/(origem.stem+'.xlsm')
        if destino.exists():raise ErroBase('A saída já existe. Escolha uma nova pasta.')
        tmp=pasta/('painel_'+uuid4().hex+'.part')
        w.save(tmp);incluir_botao(tmp,destino)
        return destino
    finally:
        w.close()
        if tmp and tmp.exists():tmp.unlink()


def main():
    p=argparse.ArgumentParser(description='Atualiza o painel preservando os atendimentos da versão 2.')
    p.add_argument('arquivo',nargs='?');p.add_argument('--saida',type=Path);p.add_argument('--sem-pausa',action='store_true')
    a=p.parse_args();code=0
    try:
        print('ATUALIZAR PAINEL | Os status e dados da Tracking serão preservados.')
        f=caminho_local(a.arquivo or input('Cole o caminho da planilha .xlsx da versão 2: '))
        print(f'Nova planilha:\n{atualizar(f,a.saida).resolve()}')
    except (KeyboardInterrupt,EOFError):code=130
    except Exception as exc:print(f'Não foi possível atualizar: {exc}');code=1
    if not a.sem_pausa:
        try:input('Pressione Enter para fechar...')
        except (KeyboardInterrupt,EOFError):pass
    return code

if __name__=='__main__':raise SystemExit(main())
