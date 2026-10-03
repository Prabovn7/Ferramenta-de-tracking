"""Abas de acompanhamento da tracking; fórmulas calculadas pelo Excel."""
from collections import Counter, defaultdict
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.formatting.rule import FormulaRule, DataBarRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.page import PageMargins

STATUS_VALIDOS = ('Pendente', 'Em andamento', 'Concluído', 'Revisar depois', 'Escalado')
BG, CARD, TEXT, MUTED, TEAL = '101923', '1C2938', 'F2F6FA', 'ADBED0', '73DEC9'
AMBER, BLUE, PURPLE = 'F5CB83', '8AAFFF', 'C6ACFA'


def texto(ws, ref, valor, size=11, color=TEXT, bold=False, fill=None, align='left'):
    if ':' in ref:
        ws.merge_cells(ref)
        cell = ws[ref.split(':')[0]]
    else:
        cell = ws[ref]
    cell.value = valor
    if isinstance(valor, str) and not valor.startswith('='):
        cell.data_type = 's'
    cell.font = Font(name='Arial', size=size, color=color, bold=bold)
    cell.alignment = Alignment(horizontal=align, vertical='center', wrap_text=True)
    if fill:
        for row in ws[ref] if ':' in ref else [[cell]]:
            for item in row:
                item.fill = PatternFill('solid', fgColor=fill)
    return cell


def literal(ws, ref, valor, **style):
    c = texto(ws, ref, valor, **style)
    c.data_type = 's'
    return c


def fundo(ws, last_row, last_col=13):
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 85
    for row in ws.iter_rows(min_row=1, max_row=last_row, max_col=last_col):
        for c in row:
            c.fill = PatternFill('solid', fgColor=BG)
            c.font = Font(name='Arial', size=11, color=TEXT)
            c.alignment = Alignment(vertical='center')
    for r in range(1, last_row+1):
        ws.row_dimensions[r].height = 23
    ws.column_dimensions['A'].width = 3
    for col in 'BCDEFGHIJKL':
        ws.column_dimensions[col].width = 12
    ws.column_dimensions['M'].width = 3
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=.25, right=.25, top=.3, bottom=.3, header=.1, footer=.1)
    ws.oddFooter.center.text = 'Tracking operacional | Página &P de &N'
    ws.oddFooter.center.size = 8
    ws.print_options.horizontalCentered = True
    ws.print_area = f'A1:M{last_row}'


def criar_painel(livro, cases, nomes, origem, aba, agora, excluidos, ausentes):
    tr = livro['Tracking']
    resumo = livro.create_sheet('Resumo')
    gestor = livro.create_sheet('Gestor', 0)
    gestor.sheet_properties.tabColor = TEAL
    tr.sheet_properties.tabColor = BLUE
    resumo.sheet_properties.tabColor = MUTED
    livro.active = 0
    livro.calculation = CalcProperties(calcId=191029, fullCalcOnLoad=True, forceFullCalc=True, calcMode='auto')

    # O resumo guarda a fotografia da distribuição inicial, inclusive exclusões.
    por_dia = defaultdict(Counter)
    totais = Counter(c.responsavel for c in cases)
    for case in cases:
        por_dia[case.abertura.date()][case.responsavel] += 1
    data_rows = [(dia, nome, ct[nome]) for dia, ct in sorted(por_dia.items()) for nome in nomes if ct[nome]]
    n = len(nomes)
    day_header = 21+n
    excl_header = day_header+len(data_rows)+3
    missing_header = excl_header+len(excluidos)+4
    end = missing_header+max(len(ausentes), 1)+4
    fundo(resumo, end)
    texto(resumo, 'B2:L3', 'Resumo da distribuição', 24, TEAL, True)
    texto(resumo, 'B4:L4', 'Registro inicial da execução. O acompanhamento do dia está na aba Gestor.', color=MUTED)
    labels = ['Base de origem', 'Aba de origem', 'Gerado em', 'Cases recebidos',
              'Excluídos: aguardo de chamado', 'Cases distribuídos', 'Pessoas informadas']
    values = [origem.name, aba, agora, len(cases)+len(excluidos), len(excluidos), len(cases), n]
    for r, label, value in zip(range(6,13), labels, values):
        texto(resumo, f'B{r}:E{r}', label, color=MUTED)
        c = texto(resumo, f'F{r}:L{r}', value, bold=True)
        if isinstance(value, str):
            c.data_type = 's'
    resumo['F8'].number_format = 'dd/mm/yyyy hh:mm:ss'
    texto(resumo, 'B14:L15', 'Mais antigos primeiro, alternando as pessoas. A alternância continua entre datas; em aberturas iguais, vale a ordem da base.', color=MUTED)
    texto(resumo, 'B17:L17', 'QUANTIDADE INICIAL POR PESSOA', bold=True, fill=CARD)
    texto(resumo, 'B18:D18', 'Responsável', bold=True)
    texto(resumo, 'E18:F18', 'Quantidade', bold=True)
    for r, nome in enumerate(nomes, 19):
        literal(resumo, f'B{r}:D{r}', nome)
        texto(resumo, f'E{r}:F{r}', totais[nome], color=TEAL, bold=True)
        resumo.row_dimensions[r].height = max(24, 16*((len(nome)+35)//36))
    livro.defined_names.add(DefinedName('Analistas', attr_text=f"'Resumo'!$B$19:$B${18+n}"))
    texto(resumo, f'B{day_header}:L{day_header}', 'DISTRIBUIÇÃO POR DATA DE ABERTURA', bold=True, fill=CARD)
    for ref, label in [(f'B{day_header+1}:D{day_header+1}', 'Data de abertura'),
                       (f'E{day_header+1}:I{day_header+1}', 'Responsável'),
                       (f'J{day_header+1}:L{day_header+1}', 'Quantidade')]:
        texto(resumo, ref, label, bold=True)
    for r, (dia, nome, qtd) in enumerate(data_rows, day_header+2):
        texto(resumo, f'B{r}:D{r}', dia).number_format = 'dd/mm/yyyy'
        literal(resumo, f'E{r}:I{r}', nome)
        texto(resumo, f'J{r}:L{r}', qtd)
    texto(resumo, f'B{excl_header}:L{excl_header}', 'EXCLUÍDOS — AGUARDO DE CHAMADO', bold=True, fill=CARD)
    for r, case in enumerate(excluidos, excl_header+1):
        literal(resumo, f'B{r}:D{r}', case.numero)
        texto(resumo, f'E{r}:H{r}', case.abertura).number_format = 'dd/mm/yyyy hh:mm:ss'
    if not excluidos:
        texto(resumo, f'B{excl_header+1}:L{excl_header+1}', 'Nenhum case excluído nesta execução.', color=MUTED)
    texto(resumo, f'B{missing_header}:L{missing_header}', 'EXCLUSÕES NÃO ENCONTRADAS NA BASE', bold=True, fill=CARD)
    for r, numero in enumerate(ausentes, missing_header+1):
        literal(resumo, f'B{r}:L{r}', numero)
    if not ausentes:
        texto(resumo, f'B{missing_header+1}:L{missing_header+1}', 'Nenhuma.', color=MUTED)
    texto(resumo, f'B{end-1}:L{end}', 'A base original foi preservada. DATA DE MODIFICAÇÃO é a data importada do Creatio; não registra as edições desta tracking.', color=MUTED)

    # A tracking continua com as sete colunas combinadas. Status por lista reduz erros.
    last = len(cases)+1
    for row in tr.iter_rows(min_row=2, max_row=last, max_col=7):
        for c in row:
            c.fill = PatternFill('solid', fgColor=CARD if c.row % 2 == 0 else BG)
            c.font = Font(name='Arial', size=11, color=TEXT if c.column in (4,5,6) else MUTED)
            c.alignment = Alignment(vertical='center', wrap_text=c.column in (5,6,7))
        tr.row_dimensions[row[0].row].height = 30
    for c in tr[1]:
        c.font = Font(name='Arial', size=11, color=AMBER if c.column in (4,5,6) else TEXT, bold=True)
        c.fill = PatternFill('solid', fgColor='263B4F')
    tr.column_dimensions['D'].width = 21
    tr.column_dimensions['G'].width = 30
    tr.sheet_view.zoomScale = 85
    status = DataValidation(type='list', formula1='"'+','.join(STATUS_VALIDOS)+'"', allow_blank=False)
    status.errorTitle = 'Escolha um status da lista'
    status.error = 'Use Pendente, Em andamento, Concluído, Revisar depois ou Escalado.'
    status.promptTitle = 'Status do case'
    status.prompt = 'Escalado continua aberto. Informe equipe/pessoa e motivo em OBSERVAÇÃO.'
    status.showErrorMessage = status.showInputMessage = True
    status.errorStyle = 'stop'
    tr.add_data_validation(status)
    status.add(f'D2:D{last}')
    owner = DataValidation(type='list', formula1='=Analistas', allow_blank=False)
    owner.showErrorMessage = True
    owner.errorTitle = 'Responsável não cadastrado'
    owner.error = 'Escolha uma das pessoas desta distribuição.'
    tr.add_data_validation(owner)
    owner.add(f'G2:G{last}')
    for label, color in zip(STATUS_VALIDOS, [MUTED, BLUE, TEAL, AMBER, PURPLE]):
        tr.conditional_formatting.add(f'D2:D{last}', FormulaRule(formula=[f'D2="{label}"'], font=Font(color=color, bold=True)))
    tr.conditional_formatting.add(f'F2:F{last}', FormulaRule(formula=['AND($D2="Escalado",LEN(TRIM($F2))=0)'], fill=PatternFill('solid', fgColor='603C23'), font=Font(color='FFE2A7')))

    from dashboard_gestor import desenhar
    desenhar(gestor, resumo, nomes, livro)
