"""Distribuidor local de cases. Python 3.10+; sem rede ou acesso de administrador."""
from __future__ import annotations

import argparse
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from math import isfinite
from pathlib import Path
import re
import sys
import unicodedata
from uuid import uuid4
from zipfile import BadZipFile

# As bibliotecas acompanham o programa; nenhum pacote e instalado no computador.
PASTA_PROGRAMA = Path(__file__).resolve().parent
sys.path.insert(0, str(PASTA_PROGRAMA))
for biblioteca in sorted((PASTA_PROGRAMA / "bibliotecas").glob("*.whl")):
    sys.path.insert(0, str(biblioteca))
try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils.datetime import from_excel
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.utils.exceptions import InvalidFileException
except ImportError:
    print("Bibliotecas ausentes. Extraia todo o ZIP, incluindo a pasta bibliotecas.")
    raise SystemExit(1)

CABECALHOS = ["CASE", "DATA DE ABERTURA", "DATA DE MODIFICAÇÃO", "STATUS",
              "MOTIVO", "OBSERVAÇÃO", "RESPONSÁVEL"]
OBRIGATORIOS = ("CASE", "DATADEABERTURA", "DATADEMODIFICACAO")
NOMES_DE_ORIGEM = {
    "CASE": "CASE",
    "NUMERODOCASO": "CASE",
    "DATADEABERTURA": "DATADEABERTURA",
    "DATADEMODIFICACAO": "DATADEMODIFICACAO",
    "MODIFICADOEM": "DATADEMODIFICACAO",
}
CONTROLES_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff\ud800-\udfff]")


class ErroBase(ValueError):
    """Problema que pode ser corrigido na planilha ou nos dados informados."""


@dataclass
class Case:
    numero: str
    abertura: datetime
    modificacao: datetime | None
    linha: int
    responsavel: str = ""


def normalizar(valor) -> str:
    texto = unicodedata.normalize("NFD", str(valor or ""))
    return re.sub(r"[^A-Z0-9]", "", "".join(c for c in texto if not unicodedata.combining(c)).upper())


def caminho_local(texto: str) -> Path:
    texto = texto.strip()
    if len(texto) >= 2 and texto[0] == texto[-1] and texto[0] in "\"'":
        texto = texto[1:-1]
    if not texto:
        raise ErroBase("Cole o caminho completo da planilha.")
    return Path(texto).expanduser()


def vazio(valor) -> bool:
    return valor is None or (isinstance(valor, str) and not valor.strip())


def validar_texto(texto: str, rotulo: str, limite: int = 32767) -> str:
    if CONTROLES_XML.search(texto) or len(texto) > limite:
        raise ErroBase(f"{rotulo}: texto muito longo ou com caracteres inválidos.")
    return texto


def ler_identificador(celula, linha: int) -> str:
    valor = celula.value
    if vazio(valor):
        raise ErroBase(f"CASE vazio na linha {linha}.")
    if isinstance(valor, bool) or isinstance(valor, (date, datetime)):
        raise ErroBase(f"CASE inválido na linha {linha}. Use o número ou código do case.")
    if isinstance(valor, (int, float)):
        if not isfinite(valor) or valor != int(valor):
            raise ErroBase(f"CASE inválido na linha {linha}.")
        texto = str(int(valor))
        # Preserva zeros iniciais quando a célula usa, por exemplo, formato 000000.
        if re.fullmatch(r"0+", celula.number_format or ""):
            texto = texto.zfill(len(celula.number_format))
    else:
        texto = str(valor).strip()
    return validar_texto(texto, f"CASE na linha {linha}")


def ler_data(valor, epoca, rotulo: str, linha: int, opcional=False) -> datetime | None:
    if vazio(valor):
        if opcional:
            return None
        raise ErroBase(f"{rotulo} vazia na linha {linha}.")
    resultado = None
    if isinstance(valor, datetime):
        resultado = valor
    elif isinstance(valor, date):
        resultado = datetime.combine(valor, datetime.min.time())
    elif isinstance(valor, (int, float)) and not isinstance(valor, bool):
        try:
            if isfinite(valor) and valor >= 1:
                resultado = from_excel(valor, epoca)
        except (ValueError, OverflowError):
            pass
    elif isinstance(valor, str):
        texto = valor.strip()
        for formato in ("%d/%m/%Y", "%d/%m/%Y %H:%M", "%d/%m/%Y %H:%M:%S",
                        "%d/%m/%Y %H:%M:%S.%f", "%Y-%m-%d", "%Y-%m-%d %H:%M",
                        "%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S",
                        "%Y-%m-%dT%H:%M:%S.%f", "%d.%m.%Y", "%d.%m.%Y %H:%M:%S"):
            try:
                resultado = datetime.strptime(texto, formato)
                break
            except ValueError:
                continue
    if not isinstance(resultado, datetime) or resultado.year < 1900 or resultado.tzinfo is not None:
        raise ErroBase(f"{rotulo} inválida na linha {linha}. Use uma data do Excel ou dd/MM/aaaa HH:mm:ss.")
    return resultado


def localizar_abas(livro, aba: str | None):
    candidatas = []
    if aba and aba not in livro.sheetnames:
        raise ErroBase(f"A aba '{aba}' não existe nessa base.")
    for planilha in livro.worksheets:
        if aba and planilha.title != aba:
            continue
        planilha.reset_dimensions()
        for numero, linha in enumerate(planilha.iter_rows(max_row=50), 1):
            mapa = {}
            repetidos = set()
            for coluna, celula in enumerate(linha):
                nome = NOMES_DE_ORIGEM.get(normalizar(celula.value))
                if nome in OBRIGATORIOS:
                    if nome in mapa:
                        repetidos.add(nome)
                    mapa[nome] = coluna
            if all(nome in mapa for nome in OBRIGATORIOS):
                if repetidos:
                    raise ErroBase(f"Cabeçalhos repetidos na aba '{planilha.title}', linha {numero}.")
                candidatas.append((planilha, numero, mapa))
                break
    if not candidatas:
        raise ErroBase("Não encontrei os três campos juntos nas primeiras 50 linhas: CASE ou Número do caso; Data de abertura; Data de modificação ou Modificado em.")
    return candidatas


def pedir_inteiro(pergunta: str, minimo: int, maximo: int) -> int:
    while True:
        try:
            valor = int(input(pergunta).strip())
            if minimo <= valor <= maximo:
                return valor
        except ValueError:
            pass
        print(f"Informe um número entre {minimo} e {maximo}.")


def carregar_base(caminho: Path, aba: str | None = None, interativo=False):
    if caminho.suffix.lower() != ".xlsx":
        raise ErroBase("Use uma planilha .xlsx. Salve a exportação nesse formato antes de continuar.")
    if not caminho.is_file():
        raise ErroBase("Arquivo não encontrado. Confira o caminho colado, incluindo o nome e .xlsx.")
    try:
        livro = load_workbook(caminho, read_only=True, data_only=False, keep_links=False)
    except (BadZipFile, InvalidFileException, KeyError, ValueError) as exc:
        raise ErroBase("Não foi possível ler esse XLSX. Salve novamente no Excel, sem senha.") from exc
    try:
        candidatas = localizar_abas(livro, aba)
        indice = 0
        if len(candidatas) > 1:
            if not interativo:
                raise ErroBase("Mais de uma aba possui os campos necessários. Informe --aba.")
            print("\nEscolha a aba da base:")
            for i, (planilha, _, _) in enumerate(candidatas, 1):
                print(f"  {i}. {planilha.title}")
            indice = pedir_inteiro("Número da aba: ", 1, len(candidatas)) - 1
        planilha, cabecalho, mapa = candidatas[indice]
        cases, encontrados = [], {}
        for numero, linha in enumerate(planilha.iter_rows(min_row=cabecalho + 1, max_col=max(mapa.values()) + 1), cabecalho + 1):
            celulas = [linha[mapa[nome]] for nome in OBRIGATORIOS]
            if all(vazio(c.value) for c in celulas):
                continue
            for celula in celulas:
                if celula.data_type in ("f", "e"):
                    raise ErroBase(f"Fórmula ou erro do Excel na linha {numero}. Use uma exportação com valores nos três campos necessários.")
            identificador = ler_identificador(celulas[0], numero)
            chave = identificador.casefold()
            if chave in encontrados:
                raise ErroBase(f"CASE duplicado nas linhas {encontrados[chave]} e {numero}. Confira a base antes de distribuir.")
            encontrados[chave] = numero
            abertura = ler_data(celulas[1].value, livro.epoch, "DATA DE ABERTURA", numero)
            modificacao = ler_data(celulas[2].value, livro.epoch, "DATA DE MODIFICAÇÃO", numero, opcional=True)
            cases.append(Case(identificador, abertura, modificacao, numero))
        if not cases:
            raise ErroBase("A aba está vazia: nenhum case foi encontrado abaixo do cabeçalho.")
        if len(cases) > 1048575:
            raise ErroBase("A quantidade de cases excede o limite de uma aba do Excel.")
        cases.sort(key=lambda case: (case.abertura, case.linha))
        return cases, planilha.title
    finally:
        livro.close()


def validar_nomes(nomes) -> list[str]:
    limpos, usados = [], set()
    for nome in nomes:
        nome = nome.strip()
        if not nome or nome.casefold() in usados:
            raise ErroBase("Os nomes precisam estar preenchidos e ser diferentes entre si.")
        validar_texto(nome, "Nome", limite=120)
        usados.add(nome.casefold())
        limpos.append(nome)
    if not limpos:
        raise ErroBase("Informe ao menos uma pessoa.")
    return limpos


def chave_case(numero: str) -> str:
    numero = str(numero).strip()
    return (numero.lstrip('0') or '0') if numero.isdigit() else numero.casefold()


def carregar_exclusoes(caminho: Path) -> list[str]:
    if not caminho.exists():
        return []
    numeros = []
    for linha, texto in enumerate(caminho.read_text(encoding='utf-8-sig').splitlines(), 1):
        texto = texto.split('#', 1)[0].strip()
        if not texto:
            continue
        for numero in re.split(r'[;,\s]+', texto):
            if not numero.isdigit():
                raise ErroBase(f'Lista de exclusões, linha {linha}: informe somente números de case.')
            numeros.append(numero)
    return list(dict.fromkeys(numeros))


def separar_exclusoes(cases: list[Case], numeros: list[str]):
    bloqueados = {chave_case(n) for n in numeros}
    encontrados = {chave_case(c.numero) for c in cases}
    elegiveis = [c for c in cases if chave_case(c.numero) not in bloqueados]
    excluidos = [c for c in cases if chave_case(c.numero) in bloqueados]
    ausentes = [n for n in numeros if chave_case(n) not in encontrados]
    return elegiveis, excluidos, ausentes


def pedir_nomes() -> list[str]:
    quantidade = pedir_inteiro("\nQuantas pessoas receberão os cases? ", 1, 1000)
    nomes = []
    while len(nomes) < quantidade:
        nome = input(f"Nome da pessoa {len(nomes) + 1}: ")
        try:
            nomes = validar_nomes(nomes + [nome])
        except ErroBase as exc:
            print(exc)
    return nomes


def distribuir(cases: list[Case], nomes: list[str]) -> None:
    nomes = validar_nomes(nomes)
    cases.sort(key=lambda case: (case.abertura, case.linha))
    # A alternância continua ao mudar o dia, para manter o total equilibrado.
    for indice, case in enumerate(cases):
        case.responsavel = nomes[indice % len(nomes)]


def salvar_resultado(cases: list[Case], nomes: list[str], origem: Path, aba: str, pasta: Path,
                     excluidos: list[Case] | None = None, ausentes: list[str] | None = None) -> Path:
    excluidos = excluidos or []
    ausentes = ausentes or []
    agora = datetime.now()
    destino = pasta / (agora.strftime("%Y-%m-%d_%H-%M-%S_") + uuid4().hex[:8])
    destino.mkdir(parents=True, exist_ok=False)
    arquivo = destino / agora.strftime("Tracking_%d-%m.xlsm")
    temporario = destino / "resultado.part"
    livro = Workbook()
    try:
        planilha = livro.active
        planilha.title = "Tracking"
        planilha.append(CABECALHOS)
        planilha.freeze_panes = "A2"
        planilha.sheet_view.showGridLines = False
        for linha, case in enumerate(cases, 2):
            planilha.append([case.numero, case.abertura, case.modificacao, "Pendente", "", "", case.responsavel])
            for coluna in (1, 4, 5, 6, 7):
                # Texto literal, mesmo quando começa com =, +, - ou @.
                planilha.cell(linha, coluna).data_type = "s"
            for coluna in (2, 3):
                planilha.cell(linha, coluna).number_format = "dd/mm/yyyy hh:mm:ss"
        tabela = Table(displayName="CasesDistribuidos", ref=f"A1:G{len(cases) + 1}")
        tabela.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        planilha.add_table(tabela)
        planilha.row_dimensions[1].height = 32
        for celula in planilha[1]:
            celula.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            celula.fill = PatternFill("solid", fgColor="16354A")
            celula.alignment = Alignment(vertical="center", wrap_text=True)
        for coluna, largura in zip("ABCDEFG", (22, 24, 24, 18, 24, 48, 28)):
            planilha.column_dimensions[coluna].width = largura
        from painel_gestor import criar_painel
        criar_painel(livro, cases, nomes, origem, aba, agora, excluidos, ausentes)
        livro.save(temporario)
        from incluir_botao_pdf import incluir_botao
        incluir_botao(temporario, arquivo)
        temporario.unlink()
        return arquivo
    except Exception:
        # Só remove arquivos temporários criados nesta execução.
        for caminho in (temporario, arquivo):
            if caminho.exists():
                caminho.unlink()
        destino.rmdir()
        raise
    finally:
        livro.close()


def main() -> int:
    parser = argparse.ArgumentParser(description="Distribui uma base de cases por data de abertura.")
    parser.add_argument("base", nargs="?", help="Caminho da base .xlsx")
    parser.add_argument("--analistas", nargs="+", help="Nomes, na ordem da distribuição")
    parser.add_argument("--aba", help="Nome da aba a importar")
    parser.add_argument(
        "--em-espera",
        "--exclusoes",
        dest="exclusoes",
        type=Path,
        default=PASTA_PROGRAMA / "Cases_em_espera.txt",
        help="Lista de cases que não devem entrar na distribuição, um número por linha",
    )
    parser.add_argument("--saida", type=Path, default=PASTA_PROGRAMA / "Saidas")
    parser.add_argument("--sem-pausa", action="store_true", help="Encerra sem pedir Enter")
    args = parser.parse_args()
    codigo = 0
    try:
        print("\nDISTRIBUIDOR DE CASES | PYTHON\nMais antigos primeiro • distribuição equilibrada\n")
        caminho_informado = args.base
        while True:
            try:
                caminho = caminho_local(caminho_informado or input("Cole o caminho completo da base .xlsx: "))
                print("Lendo e conferindo a base...")
                cases, aba = carregar_base(caminho, args.aba, interativo=not args.sem_pausa)
                break
            except (ErroBase, OSError) as exc:
                if args.base or args.sem_pausa:
                    raise
                print(f"\nNão consegui abrir a base: {exc}\nCole outro caminho ou pressione Ctrl+C para sair.\n")
                caminho_informado = None
        print(f"\nAba: {aba} | Total recebido: {len(cases)} cases")
        cases, excluidos, ausentes = separar_exclusoes(cases, carregar_exclusoes(args.exclusoes))
        print(
            f"Cases em espera: {len(excluidos)} | "
            f"Para distribuir: {len(cases)}"
        )
        if ausentes:
            print(f"Aviso: {len(ausentes)} número(s) da lista de exclusões não aparecem nesta base.")
        if not cases:
            raise ErroBase("Todos os cases desta base estão na lista de exclusões. Nenhuma distribuição foi gerada.")
        print(f"Aberturas: {cases[0].abertura:%d/%m/%Y} a {cases[-1].abertura:%d/%m/%Y}")
        nomes = validar_nomes(args.analistas) if args.analistas else pedir_nomes()
        distribuir(cases, nomes)
        arquivo = salvar_resultado(cases, nomes, caminho, aba, args.saida, excluidos, ausentes)
        print("\nDISTRIBUIÇÃO CONCLUÍDA")
        totais = Counter(case.responsavel for case in cases)
        for nome in nomes:
            print(f"  {nome}: {totais[nome]} cases")
        print(f"\nExcel gerado:\n{arquivo.resolve()}")
        print("\nA planilha contém as abas Gestor, Tracking, Resumo e SLA.")
        print("Atualize o STATUS pela lista da aba Tracking; o painel Gestor acompanha.")
        print("Para gerar o PDF, abra no Excel instalado e clique em GERAR RESUMO PDF.")
    except (KeyboardInterrupt, EOFError):
        print("\nOperação cancelada.")
        codigo = 130
    except Exception as exc:
        print(f"\nNão foi possível concluir: {exc}")
        print("Confira a base e se você pode gravar na pasta do programa.")
        codigo = 1
    if not args.sem_pausa:
        try:
            input("\nPressione Enter para fechar...")
        except (KeyboardInterrupt, EOFError):
            pass
    return codigo


if __name__ == "__main__":
    raise SystemExit(main())
