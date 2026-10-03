"""Painel operacional com filtro, indicadores e prioridades por fórmulas."""
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.pagebreak import Break
from painel_gestor import texto, fundo, BG, CARD, TEXT, MUTED, TEAL, AMBER, BLUE, PURPLE


def desenhar(gestor, resumo, nomes, livro):
    n=len(nomes)
    sla_inicio=max(29,22+n)
    bottom=sla_inicio+15
    fundo(gestor,bottom,23)
    for c in range(2,23):gestor.column_dimensions[get_column_letter(c)].width=8.3
    gestor.column_dimensions['W'].width=3
    for r in range(1,bottom+1):gestor.row_dimensions[r].height=20
    gestor.sheet_view.zoomScale=90
    gestor.sheet_view.showRowColHeaders=False
    gestor.print_area=f'A1:W{bottom}'
    gestor.page_setup.fitToHeight=0
    gestor.print_title_rows=None
    gestor.row_breaks.append(Break(id=sla_inicio-1))
    texto(gestor,'B2:O3','Visão da operação',26,TEXT,True)
    # Área reservada ao botão nativo (inserido no XLSM).
    texto(gestor,'R2:V3','',12,BG,True,BG,align='center')
    texto(gestor,'B4:V4','TRACKING DE CASES  /  Acompanhamento do trabalho da equipe',10,TEAL,True)
    texto(gestor,'B5:D5','Visualizar:',11,MUTED)
    texto(gestor,'E5:J5','Toda a equipe',12,TEXT,True,'324459')
    texto(gestor,'L5:V5','Escolha a equipe inteira ou um analista na lista.',10,MUTED)
    # Lista em colunas auxiliares da aba Resumo, sem duplicar indicadores.
    resumo['N1']='Filtro do painel'
    resumo['N2']='Toda a equipe'
    for r,nome in enumerate(nomes,3):
        resumo.cell(r,14,nome).data_type='s'
    resumo.column_dimensions['N'].hidden=True
    livro.defined_names.add(DefinedName('FiltroGestor',attr_text=f"'Resumo'!$N$2:$N${n+2}"))
    validation=DataValidation(type='list',formula1='=FiltroGestor',allow_blank=False)
    validation.errorTitle='Escolha uma opção da lista'
    validation.error='Selecione Toda a equipe ou um analista cadastrado.'
    validation.showErrorMessage=True
    gestor.add_data_validation(validation);validation.add('E5')
    ids='CasesDistribuidos[CASE]';sts='CasesDistribuidos[STATUS]';owners='CasesDistribuidos[RESPONSÁVEL]'
    criterio='SUBSTITUTE(SUBSTITUTE(SUBSTITUTE($E$5,"~","~~"),"*","~*"),"?","~?")'
    def count(status=None,extra=''):
        filters=f'{ids},"<>"'+(f',{sts},"{status}"' if status else '')+extra
        return f'=IF($E$5="Toda a equipe",COUNTIFS({filters}),COUNTIFS({filters},{owners},{criterio}))'
    cards=[('B7:G7','B8:G10','DISTRIBUÍDOS',count(),BLUE),
           ('I7:N7','I8:N10','CONCLUÍDOS',count('Concluído'),TEAL),
           ('P7:V7','P8:V10','AINDA ABERTOS','=B8-I8',AMBER)]
    for labelref,valref,label,formula,color in cards:
        texto(gestor,labelref,label,11,MUTED,True,CARD)
        texto(gestor,valref,formula,34,color,True,CARD).number_format='#,##0'
    for labelref,valref,label,status,color in [
        ('B12:E12','B13:E14','Para iniciar','Pendente',MUTED),
        ('G12:J12','G13:J14','Em andamento','Em andamento',BLUE),
        ('L12:O12','L13:O14','Revisar depois','Revisar depois',AMBER),
        ('Q12:V12','Q13:V14','Escalados','Escalado',PURPLE)]:
        texto(gestor,labelref,label,11,MUTED,False,CARD)
        texto(gestor,valref,count(status),24,color,True,CARD)
    texto(gestor,'B16:E16','CONCLUSÃO',11,MUTED,True)
    texto(gestor,'F16:H16','=IF(B8=0,0,I8/B8)',16,TEAL,True).number_format='0.0%'
    for c in range(9,23):
        cell=f'{get_column_letter(c)}16'
        gestor[cell].fill=PatternFill('solid',fgColor=CARD)
        gestor.conditional_formatting.add(cell,FormulaRule(formula=[f'$F$16>=({c}-8)/14'],fill=PatternFill('solid',fgColor='FF'+TEAL,bgColor='FF'+TEAL)))
    texto(gestor,'B17:V17','Abertos incluem pendentes, em andamento, revisão e escalados. Escalado não conta como concluído.',10,MUTED)
    texto(gestor,'B19:M19','DISTRIBUIÇÃO E PROGRESSO DA EQUIPE',11,TEAL,True)
    texto(gestor,'O19:V19','ATENÇÃO AGORA',11,AMBER,True)
    for ref,label in [('B20:E20','Responsável'),('F20:G20','Total'),('H20:I20','Concluídos'),('J20:K20','Abertos'),('L20:M20','Conclusão')]:
        texto(gestor,ref,label,10,MUTED,True,CARD)
    for r,nome in enumerate(nomes,21):
        texto(gestor,f'B{r}:E{r}',f"='Resumo'!B{r-2}")
        cr=f'SUBSTITUTE(SUBSTITUTE(SUBSTITUTE($B{r},"~","~~"),"*","~*"),"?","~?")'
        texto(gestor,f'F{r}:G{r}',f'=COUNTIFS({ids},"<>",{owners},{cr})',12,BLUE,True,align='center')
        texto(gestor,f'H{r}:I{r}',f'=COUNTIFS({ids},"<>",{owners},{cr},{sts},"Concluído")',12,TEAL,True,align='center')
        texto(gestor,f'J{r}:K{r}',f'=F{r}-H{r}',12,AMBER,True,align='center')
        texto(gestor,f'L{r}:M{r}',f'=IF(F{r}=0,0,H{r}/F{r})',11,TEXT,align='center').number_format='0.0%'
        gestor.row_dimensions[r].height=max(26,15*((len(nome)+34)//35))
        gestor.conditional_formatting.add(f'B{r}:M{r}',FormulaRule(formula=[f'AND($E$5<>"Toda a equipe",$B{r}=$E$5)'],fill=PatternFill('solid',fgColor='FF304356',bgColor='FF304356')))
    texto(gestor,'O20:V20','ABERTURA MAIS ANTIGA EM ABERTO',9,MUTED)
    date='CasesDistribuidos[DATA DE ABERTURA]'
    filters=f'{date},{ids},"<>",{sts},"<>Concluído"'
    formula=f'=IF(P8=0,"Sem cases abertos",IF($E$5="Toda a equipe",_xlfn.MINIFS({filters}),_xlfn.MINIFS({filters},{owners},{criterio})))'
    texto(gestor,'O21:V22',formula,18,AMBER,True,CARD).number_format='dd/mm/yyyy'
    texto(gestor,'O23:T23','Status vazio / fora da lista',10,MUTED)
    texto(gestor,'U23:V23','=B8-SUM(I8,B13,G13,L13,Q13)',12,AMBER,True)
    texto(gestor,'O24:T24','Escalados sem observação',10,MUTED)
    blank=f'({sts}="Escalado")*(LEN(TRIM(CasesDistribuidos[OBSERVAÇÃO]))=0)*({ids}<>"")'
    texto(gestor,'U24:V24',f'=IF($E$5="Toda a equipe",SUMPRODUCT({blank}),SUMPRODUCT({blank}*({owners}=$E$5)))',12,AMBER,True)
    texto(gestor,'O25:T25','Sem responsável cadastrado¹',10,MUTED)
    texto(gestor,'U25:V25',f'=COUNTA({ids})-SUM(F21:F{20+n})',12,AMBER,True)
    texto(gestor,'O26:V27','¹ Alerta da equipe inteira. Corrija os campos indicados na Tracking.',9,MUTED)
    from sla_cases import acrescentar_sla
    acrescentar_sla(livro,gestor,sla_inicio)
    nav=sla_inicio+12
    c=texto(gestor,f'B{nav}:F{nav+1}','ABRIR TRACKING  →',11,TEAL,True,CARD,align='center');c.hyperlink="#'Tracking'!A1"
    c=texto(gestor,f'H{nav}:M{nav+1}','DISTRIBUIÇÃO INICIAL  →',11,TEAL,True,CARD,align='center');c.hyperlink="#'Resumo'!B2"
    texto(gestor,f'O{nav}:V{nav+1}','PDF = visão atual do painel, incluindo o filtro selecionado.',10,MUTED)
    texto(gestor,f'B{bottom}:V{bottom}','Fórmulas atualizam os indicadores ao editar a Tracking. Para gerar PDF, abra no Excel instalado.',9,MUTED)
